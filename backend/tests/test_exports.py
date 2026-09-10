from __future__ import annotations

import hashlib
import io
import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from enterprise_insight_backend.exporting import ExportService


def _project_with_projection(client: TestClient) -> str:
    company = client.post(
        "/api/v3/companies", json={"name": "导出测试企业", "industry": "制造业"}
    ).json()
    project = client.post(
        f"/api/v3/companies/{company['id']}/projects", json={"name": "企业建模"}
    ).json()
    project_id = project["id"]
    assert client.post(f"/api/v3/projects/{project_id}/ontology/default-pack").status_code == 200
    entity = client.post(
        f"/api/v3/projects/{project_id}/entities",
        json={
            "type_key": "organization_unit",
            "stable_key": "department.operations",
            "name": "运营部",
            "properties": {"mandate": "交付客户价值"},
            "viewpoint": "DESIGNED",
            "evidence": [],
        },
    )
    assert entity.status_code == 201, entity.text
    return project_id


def _create_export(client: TestClient, project_id: str, export_format: str) -> dict:
    response = client.post(
        f"/api/v3/projects/{project_id}/exports",
        json={
            "format": export_format,
            "include_evidence": True,
            "include_lineage": True,
        },
    )
    assert response.status_code == 201, response.text
    payload = response.json()
    assert payload["status"] == "COMPLETED", payload
    assert payload["download_url"]
    return payload


def test_json_and_csv_exports_are_generated_and_downloadable(client: TestClient) -> None:
    project_id = _project_with_projection(client)

    json_job = _create_export(client, project_id, "json")
    json_response = client.get(json_job["download_url"])
    assert json_response.status_code == 200
    bundle = json.loads(json_response.content.decode("utf-8"))
    assert bundle["company"]["name"] == "导出测试企业"
    assert bundle["graph"]["entities"][0]["name"] == "运营部"
    assert "connection_profile" not in json_response.text

    csv_job = _create_export(client, project_id, "csv")
    csv_response = client.get(csv_job["download_url"])
    assert csv_response.status_code == 200
    csv_text = csv_response.content.decode("utf-8-sig")
    assert "record_kind,id,type_key,name,status,payload" in csv_text
    assert "运营部" in csv_text

    listing = client.get(f"/api/v3/projects/{project_id}/exports")
    assert listing.status_code == 200
    assert listing.json()["total"] == 2
    detail = client.get(f"/api/v3/projects/{project_id}/exports/{json_job['id']}")
    assert detail.status_code == 200
    assert detail.json()["id"] == json_job["id"]


def test_xlsx_and_bundle_exports_contain_expected_artifacts(client: TestClient) -> None:
    project_id = _project_with_projection(client)

    xlsx_job = _create_export(client, project_id, "xlsx")
    xlsx_response = client.get(xlsx_job["download_url"])
    workbook = load_workbook(io.BytesIO(xlsx_response.content), read_only=True)
    assert {"项目", "对象", "关系", "事件", "管理假设", "管理方案"}.issubset(workbook.sheetnames)
    assert any(cell.value == "运营部" for row in workbook["对象"].iter_rows() for cell in row)

    bundle_job = _create_export(client, project_id, "bundle")
    bundle_response = client.get(bundle_job["download_url"])
    with ZipFile(io.BytesIO(bundle_response.content)) as archive:
        assert set(archive.namelist()) == {
            "enterprise.json",
            "records.csv",
            "enterprise.xlsx",
            "restore.json",
            "restore.manifest.json",
        }
        exported = json.loads(archive.read("enterprise.json"))
        assert exported["project"]["name"] == "企业建模"
        restore_bytes = archive.read("restore.json")
        restore = json.loads(restore_bytes)
        manifest = json.loads(archive.read("restore.manifest.json"))
        assert restore["schema"] == "enterprise-insight.project-archive.v1"
        assert restore["tables"]["entities"][0]["name"] == "运营部"
        assert manifest["payload_sha256"] == hashlib.sha256(restore_bytes).hexdigest()


def test_export_rejects_unknown_format(client: TestClient) -> None:
    project_id = _project_with_projection(client)
    response = client.post(f"/api/v3/projects/{project_id}/exports", json={"format": "parquet"})
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "REQUEST_VALIDATION_FAILED"


def test_export_write_failure_preserves_existing_file_and_cleans_temp(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    destination = tmp_path / "exports" / "job.json"
    destination.parent.mkdir()
    destination.write_text("previous", encoding="utf-8")
    original_write_text = Path.write_text

    def partial_write(path: Path, data: str, *args: object, **kwargs: object) -> int:
        original_write_text(path, "partial", *args, **kwargs)
        raise OSError("simulated disk-full")

    monkeypatch.setattr(Path, "write_text", partial_write)

    with pytest.raises(OSError, match="disk-full"):
        ExportService.__new__(ExportService)._write(destination, "json", {"value": 1})

    assert destination.read_text(encoding="utf-8") == "previous"
    assert list(destination.parent.glob(f".{destination.name}.*.tmp")) == []
