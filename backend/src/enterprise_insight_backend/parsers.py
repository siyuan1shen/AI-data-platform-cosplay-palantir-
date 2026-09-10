from __future__ import annotations

import csv
import io
import json
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from openpyxl import load_workbook  # type: ignore[import-untyped]

from enterprise_insight_backend.errors import DomainError
from enterprise_insight_backend.schemas import ImportKind

MAX_IMPORT_BYTES = 25 * 1024 * 1024
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_ARCHIVE_MEMBER_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_XLSX_ROWS = 100_000
MAX_XLSX_COLUMNS = 500
MAX_PDF_PAGES = 500
WORD_NAMESPACE = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


@dataclass(frozen=True)
class ParsedDocument:
    encoding: str | None
    columns: list[str]
    rows: list[dict[str, Any]]
    warnings: list[str]
    parser_version: str = "v1"

    @property
    def document_like(self) -> bool:
        return not self.columns and all("text" in row for row in self.rows)


def parse_import(file_name: str, content: bytes, kind: ImportKind) -> ParsedDocument:
    if not content:
        raise DomainError("IMPORT_EMPTY_FILE", "文件为空，无法解析。", status_code=422)
    if len(content) > MAX_IMPORT_BYTES:
        raise DomainError(
            "IMPORT_FILE_TOO_LARGE",
            "单个材料暂时不能超过 25MB。",
            status_code=413,
        )
    _validate_extension(file_name, kind)
    try:
        if kind in {ImportKind.CSV, ImportKind.COMPANYCHECK_CSV}:
            return _parse_csv(content)
        if kind == ImportKind.TXT:
            return _parse_text(content)
        if kind == ImportKind.JSON:
            return _parse_json(content)
        if kind == ImportKind.XLSX:
            return _parse_xlsx(content)
        if kind == ImportKind.DOCX:
            return _parse_docx(content)
        if kind == ImportKind.PDF:
            return _parse_pdf(content)
    except DomainError:
        raise
    except (
        csv.Error,
        ElementTree.ParseError,
        json.JSONDecodeError,
        OSError,
        ValueError,
        zipfile.BadZipFile,
    ) as exc:
        raise DomainError(
            "IMPORT_PARSE_FAILED",
            f"{kind.value} 文件无法解析，请检查文件是否损坏或类型是否选择正确。",
            status_code=422,
            details=[{"exception": type(exc).__name__}],
        ) from exc
    except Exception as exc:  # pragma: no cover - third-party parser boundary
        raise DomainError(
            "IMPORT_PARSE_FAILED",
            f"{kind.value} 文件无法解析，请检查文件是否损坏或类型是否选择正确。",
            status_code=422,
            details=[{"exception": type(exc).__name__}],
        ) from exc
    raise DomainError("IMPORT_KIND_NOT_SUPPORTED", "该材料类型暂不支持。", status_code=422)


def _parse_csv(content: bytes) -> ParsedDocument:
    text, encoding = decode_text(content)
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    columns = [item.strip() for item in (reader.fieldnames or []) if item and item.strip()]
    if not columns:
        raise DomainError("IMPORT_NO_COLUMNS", "CSV中没有可识别的表头。", status_code=422)
    rows = [
        {
            str(key).strip(): (value or "").strip()
            for key, value in row.items()
            if key is not None
        }
        for row in reader
    ]
    return _require_rows(ParsedDocument(encoding, columns, rows, []), "CSV中没有数据行。")


def _parse_text(content: bytes) -> ParsedDocument:
    text, encoding = decode_text(content)
    paragraphs = [item.strip() for item in text.replace("\r\n", "\n").split("\n\n") if item.strip()]
    if not paragraphs and text.strip():
        paragraphs = [text.strip()]
    rows = [
        {"locator": f"paragraph:{index}", "text": paragraph}
        for index, paragraph in enumerate(paragraphs, start=1)
    ]
    return _require_rows(ParsedDocument(encoding, [], rows, []), "文本文件没有可读取内容。")


def _parse_json(content: bytes) -> ParsedDocument:
    text, encoding = decode_text(content)
    value = json.loads(text)
    if isinstance(value, dict):
        items = [value]
    elif isinstance(value, list) and all(isinstance(item, dict) for item in value):
        items = value
    else:
        raise DomainError(
            "IMPORT_JSON_SHAPE_UNSUPPORTED",
            "JSON 顶层必须是对象，或由对象组成的数组。",
            status_code=422,
        )
    columns = list(dict.fromkeys(str(key) for item in items for key in item))
    rows = [
        {
            **{str(key): _json_value(value) for key, value in item.items()},
            "__locator__": f"json:item:{index}",
        }
        for index, item in enumerate(items, start=1)
    ]
    return _require_rows(ParsedDocument(encoding, columns, rows, []), "JSON 中没有记录。")


def _parse_xlsx(content: bytes) -> ParsedDocument:
    _validate_zip_limits(content)
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    rows: list[dict[str, Any]] = []
    columns: list[str] = []
    warnings: list[str] = []
    try:
        for worksheet in workbook.worksheets:
            values = worksheet.iter_rows(values_only=True)
            header_values = next(values, None)
            if header_values is None:
                continue
            headers = [str(value).strip() if value is not None else "" for value in header_values]
            if len(headers) > MAX_XLSX_COLUMNS:
                raise DomainError(
                    "IMPORT_WORKBOOK_TOO_LARGE",
                    f"工作表“{worksheet.title}”列数超过 {MAX_XLSX_COLUMNS} 列上限。",
                    status_code=413,
                )
            if not any(headers):
                warnings.append(f"工作表“{worksheet.title}”没有表头，已跳过。")
                continue
            seen_headers: set[str] = set()
            for header in headers:
                if not header:
                    continue
                if header in seen_headers:
                    raise DomainError(
                        "IMPORT_DUPLICATE_COLUMNS",
                        f"工作表“{worksheet.title}”存在重复列名：{header}。",
                        status_code=422,
                    )
                seen_headers.add(header)
                if header not in columns:
                    columns.append(header)
            for row_number, values_row in enumerate(values, start=2):
                if row_number > MAX_XLSX_ROWS + 1:
                    raise DomainError(
                        "IMPORT_WORKBOOK_TOO_LARGE",
                        f"工作表“{worksheet.title}”行数超过 {MAX_XLSX_ROWS} 行上限。",
                        status_code=413,
                    )
                record = {
                    header: _json_value(value)
                    for header, value in zip(headers, values_row, strict=False)
                    if header
                }
                if not any(value not in (None, "") for value in record.values()):
                    continue
                record["__locator__"] = f"sheet:{worksheet.title}:row:{row_number}"
                rows.append(record)
    finally:
        workbook.close()
    return _require_rows(
        ParsedDocument(None, columns, rows, warnings),
        "Excel 工作簿中没有可读取的数据行。",
    )


def _parse_docx(content: bytes) -> ParsedDocument:
    _validate_zip_limits(content)
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        try:
            document_xml = _read_zip_member_limited(archive, "word/document.xml")
        except KeyError as exc:
            raise DomainError(
                "IMPORT_DOCX_STRUCTURE_INVALID",
                "Word 文件缺少正文结构。",
                status_code=422,
            ) from exc
    root = ElementTree.fromstring(document_xml)
    body = root.find(f"{{{WORD_NAMESPACE}}}body")
    if body is None:
        raise DomainError("IMPORT_DOCX_STRUCTURE_INVALID", "Word 文件没有正文。", status_code=422)
    rows: list[dict[str, Any]] = []
    paragraph_index = 0
    table_index = 0
    for child in body:
        if child.tag == f"{{{WORD_NAMESPACE}}}p":
            text = _word_text(child)
            if text:
                paragraph_index += 1
                rows.append({"locator": f"paragraph:{paragraph_index}", "text": text})
        elif child.tag == f"{{{WORD_NAMESPACE}}}tbl":
            table_index += 1
            for row_index, table_row in enumerate(
                child.findall(f"{{{WORD_NAMESPACE}}}tr"), start=1
            ):
                cells = [_word_text(cell) for cell in table_row.findall(f"{{{WORD_NAMESPACE}}}tc")]
                if any(cells):
                    rows.append(
                        {
                            "locator": f"table:{table_index}:row:{row_index}",
                            "text": " | ".join(cells),
                        }
                    )
    return _require_rows(ParsedDocument(None, [], rows, []), "Word 文件没有可读取的正文或表格。")


def _parse_pdf(content: bytes) -> ParsedDocument:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - exercised in minimal installations
        raise DomainError(
            "IMPORT_PARSER_UNAVAILABLE",
            "PDF 解析组件尚未安装，请运行项目依赖安装后重试。",
            status_code=503,
        ) from exc
    reader = PdfReader(io.BytesIO(content))
    if len(reader.pages) > MAX_PDF_PAGES:
        raise DomainError(
            "IMPORT_PDF_TOO_LARGE",
            f"PDF 页数超过 {MAX_PDF_PAGES} 页上限。",
            status_code=413,
        )
    rows: list[dict[str, Any]] = []
    empty_pages: list[int] = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            rows.append({"locator": f"page:{page_number}", "text": text})
        else:
            empty_pages.append(page_number)
    if not rows:
        raise DomainError(
            "OCR_REQUIRED",
            "PDF 没有可提取文字，可能是扫描件，需要先进行 OCR。",
            status_code=422,
        )
    warnings = []
    if empty_pages:
        warnings.append(f"第 {', '.join(map(str, empty_pages))} 页没有提取到文字，可能需要 OCR。")
    return ParsedDocument(None, [], rows, warnings)


def decode_text(content: bytes) -> tuple[str, str]:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return content.decode(encoding), encoding
        except UnicodeDecodeError:
            continue
    raise DomainError("IMPORT_ENCODING_UNKNOWN", "无法识别文件字符编码。", status_code=422)


def _validate_extension(file_name: str, kind: ImportKind) -> None:
    suffix = Path(file_name).suffix.lower()
    allowed = {
        ImportKind.COMPANYCHECK_CSV: {".csv"},
        ImportKind.CSV: {".csv", ".tsv"},
        ImportKind.XLSX: {".xlsx"},
        ImportKind.DOCX: {".docx"},
        ImportKind.PDF: {".pdf"},
        ImportKind.TXT: {".txt", ".md", ".markdown"},
        ImportKind.JSON: {".json"},
    }[kind]
    if suffix not in allowed:
        raise DomainError(
            "IMPORT_FILE_EXTENSION_MISMATCH",
            f"文件扩展名 {suffix or '（无）'} 与选择的 {kind.value} 类型不匹配。",
            status_code=422,
        )


def _require_rows(parsed: ParsedDocument, message: str) -> ParsedDocument:
    if not parsed.rows:
        raise DomainError("IMPORT_NO_ROWS", message, status_code=422)
    return parsed


def _validate_zip_limits(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            total_size = 0
            for info in archive.infolist():
                if info.file_size > MAX_ARCHIVE_MEMBER_UNCOMPRESSED_BYTES:
                    raise DomainError(
                        "IMPORT_ARCHIVE_MEMBER_TOO_LARGE",
                        "压缩材料解压后的单个文件超过大小上限。",
                        status_code=413,
                    )
                total_size += info.file_size
                if total_size > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
                    raise DomainError(
                        "IMPORT_ARCHIVE_TOO_LARGE",
                        "压缩材料解压后的总内容超过大小上限。",
                        status_code=413,
                    )
    except DomainError:
        raise
    except (OSError, zipfile.BadZipFile) as exc:
        raise DomainError(
            "IMPORT_PARSE_FAILED",
            "压缩材料无法读取，请检查文件是否损坏。",
            status_code=422,
            details=[{"exception": type(exc).__name__}],
        ) from exc


def _read_zip_member_limited(archive: zipfile.ZipFile, name: str) -> bytes:
    info = archive.getinfo(name)
    if info.file_size > MAX_ARCHIVE_MEMBER_UNCOMPRESSED_BYTES:
        raise DomainError(
            "IMPORT_ARCHIVE_MEMBER_TOO_LARGE",
            "压缩材料解压后的正文超过大小上限。",
            status_code=413,
        )
    return archive.read(name)


def _word_text(element: ElementTree.Element) -> str:
    return "".join(
        node.text or "" for node in element.iter(f"{{{WORD_NAMESPACE}}}t")
    ).strip()


def _json_value(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value
